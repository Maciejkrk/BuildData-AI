import json
from io import BytesIO

import pytest
from openpyxl import Workbook

from data_master_app.converter import analyze_uploaded_file, analyze_product_model_files, convert_products_file, normalize_header


def bundle():
    return {
        'productsModels.json': json.dumps({'models': [
            {'Id': 5010, 'Name': 'Product', 'modelType': 'Product'},
            {'Id': 5020, 'Name': 'Documents', 'modelType': 'Attribute'},
        ]}),
        'productsAttributes.json': json.dumps({'attributes': [
            {'Id': 5100, 'ProductModelId': 5010, 'DispName': 'Name', 'AttributeType': 'Text'},
            {'Id': 5101, 'ProductModelId': 5010, 'DispName': 'Documents', 'AttributeType': 'Model_Array', 'TargetModelId': 5020},
            {'Id': 5102, 'ProductModelId': 5020, 'DispName': 'Document title', 'AttributeType': 'Text'},
            {'Id': 5103, 'ProductModelId': 5020, 'DispName': 'Language', 'AttributeType': 'Text'},
        ]})
    }


def source():
    return json.dumps({
        'Products': [{'Name': 'Same name', 'Key': 'P1'}, {'Name': 'Same name', 'Key': 'P2'}, {'Name': 'No documents', 'Key': ''}],
        'Documents': [
            {'Owner': 'P1', 'Title': 'Technical sheet', 'Lang': 'pl'},
            {'Owner': ' P1 ', 'Title': 'Technical sheet', 'Lang': 'en'},
            {'Owner': 'P2', 'Title': 'Certificate', 'Lang': 'pl'},
            {'Owner': '', 'Title': 'Unlinked', 'Lang': 'pl'},
            {'Owner': 'unknown', 'Title': 'Other', 'Lang': 'pl'},
        ]
    }).encode()


def profile():
    return {'Name': {'target_path': 'product.name.value'}, '_product_table': 'Products', '_nested_relations': {
        '5101': {'enabled': True, 'table': 'Documents', 'product_key': 'Key', 'child_key': 'Owner', 'fields': {'5102': 'Title', '5103': 'Lang'}}
    }}


@pytest.mark.parametrize('relation_type', ['Model_Array', 'Table_Model'])
def test_lists_join_all_records_and_keep_row_identity_after_profile_roundtrip(tmp_path, relation_type):
    files = bundle()
    files['productsAttributes.json'] = files['productsAttributes.json'].replace('Model_Array', relation_type)
    analysis = analyze_uploaded_file('source.json', source(), product_model_files=files)
    assert analysis['nested_models'][0]['label'] == 'Documents'
    assert len(analysis['nested_models'][0]['fields']) == 2
    saved = json.loads(json.dumps(profile()))
    result = convert_products_file('source.json', source(), tmp_path, product_model_files=files, product_mapping_profile=saved)
    assert result['report']['nested_relations'][0]['matched_records'] == 3
    assert result['report']['nested_relations'][0]['unmatched_records'] == 1
    payload = json.loads((tmp_path / result['job_id'] / 'products.json').read_text(encoding='utf-8'))
    assert len(payload['products']) == 3
    attrs = [[a for a in p['dataVersions'][0]['productAttributes'] if a['ParentAttributeId'] == 5101] for p in payload['products']]
    assert [len(items) for items in attrs] == [4, 2, 0]
    assert {a['RowI'] for a in attrs[0]} == {1, 2}
    assert len({a['hash'] for a in attrs[0]}) == 2
    assert {a['MainAttributeId'] for a in attrs[0]} == {5020}
    assert {a['TextValue'] for a in attrs[0] if a['AttributeId'] == 5103} == {'pl', 'en'}


@pytest.mark.parametrize('field,value', [('table', 'Missing'), ('child_key', 'Missing'), ('product_key', 'Missing'), ('fields', {})])
def test_incomplete_join_is_rejected_instead_of_silently_exporting_empty_lists(tmp_path, field, value):
    config = profile()
    config['_nested_relations']['5101'][field] = value
    with pytest.raises(ValueError):
        convert_products_file('source.json', source(), tmp_path, product_model_files=bundle(), product_mapping_profile=config)


def test_excel_worksheets_are_joined_using_saved_mapping(tmp_path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in json.loads(source()).items():
        sheet = workbook.create_sheet(name)
        columns = list(rows[0])
        sheet.append(columns)
        for row in rows:
            sheet.append([row[column] for column in columns])
    stream = BytesIO()
    workbook.save(stream)
    config = profile()
    config[normalize_header('Name')] = config.pop('Name')
    relation = config['_nested_relations']['5101']
    relation['product_key'] = normalize_header(relation['product_key'])
    relation['child_key'] = normalize_header(relation['child_key'])
    relation['fields'] = {key: normalize_header(column) for key, column in relation['fields'].items()}
    result = convert_products_file('source.xlsx', stream.getvalue(), tmp_path, product_model_files=bundle(), product_mapping_profile=config)
    assert result['report']['nested_relations'][0]['matched_records'] == 3


def test_model_loading_exposes_repeatable_fields():
    result = analyze_product_model_files(bundle())
    assert result['nested_models'][0]['parent_attribute_id'] == 5101
    assert len(result['nested_models'][0]['fields']) == 2


def test_single_model_source_rejects_conflicting_objects(tmp_path):
    files = bundle()
    files['productsAttributes.json'] = files['productsAttributes.json'].replace('Model_Array', 'Model')
    assert analyze_product_model_files(files)['nested_models'][0]['multiple'] is False
    with pytest.raises(ValueError, match='Multiple objects for single model'):
        convert_products_file('source.json', source(), tmp_path, product_model_files=files, product_mapping_profile=profile())


def test_record_key_deduplicates_identical_rows_and_rejects_conflicts(tmp_path):
    data = json.loads(source())
    data['Documents'] = [{'Owner': 'P1', 'Title': 'A', 'Lang': 'pl'}] * 2
    config = profile()
    config['_nested_relations']['5101']['record_key'] = 'Title'
    result = convert_products_file('source.json', json.dumps(data).encode(), tmp_path, product_model_files=bundle(), product_mapping_profile=config)
    assert result['report']['nested_relations'][0]['matched_records'] == 1
    data['Documents'][1] = {'Owner': 'P1', 'Title': 'A', 'Lang': 'en'}
    with pytest.raises(ValueError, match='Conflicting record key'):
        convert_products_file('source.json', json.dumps(data).encode(), tmp_path, product_model_files=bundle(), product_mapping_profile=config)


def test_primary_product_key_preserves_same_named_products(tmp_path):
    data = json.loads(source())
    data['Products'] = data['Products'][:2]
    config = {'Name': {'target_path': 'product.name.value'}, '_product_table': 'Products', '_product_key': 'Key'}
    result = convert_products_file('source.json', json.dumps(data).encode(), tmp_path, product_model_files=bundle(), product_mapping_profile=config)
    payload = json.loads((tmp_path / result['job_id'] / 'products.json').read_text(encoding='utf-8'))
    assert len(payload['products']) == 2
    config['_product_key'] = 'Missing'
    with pytest.raises(ValueError, match='Missing product key'):
        convert_products_file('source.json', json.dumps(data).encode(), tmp_path, product_model_files=bundle(), product_mapping_profile=config)


def test_ambiguous_parent_key_is_rejected(tmp_path):
    config = profile()
    config['_product_key'] = 'Key'
    config['_nested_relations']['5101']['product_key'] = 'Name'
    with pytest.raises(ValueError, match='Ambiguous product join key'):
        convert_products_file('source.json', source(), tmp_path, product_model_files=bundle(), product_mapping_profile=config)


def test_external_relation_uses_primary_identifier_not_a_separate_parent_key(tmp_path):
    data = json.loads(source())
    data['Products'] = data['Products'][:2]
    config = profile()
    config['_product_key'] = 'Key'
    config['_nested_relations']['5101'].update(use_primary_key=True, product_key='Name')
    result = convert_products_file('source.json', json.dumps(data).encode(), tmp_path, product_model_files=bundle(), product_mapping_profile=config)
    assert result['report']['nested_relations'][0]['matched_records'] == 3
    del config['_product_key']
    with pytest.raises(ValueError, match='primary product identifier'):
        convert_products_file('source.json', json.dumps(data).encode(), tmp_path, product_model_files=bundle(), product_mapping_profile=config)


def test_same_worksheet_can_supply_product_and_multiple_child_records(tmp_path):
    rows = [{'Name': 'A', 'Key': 'P1', 'Title': 'One'}, {'Name': 'A', 'Key': 'P1', 'Title': 'Two'}]
    config = profile()
    config['_product_key'] = 'Key'
    config['_nested_relations']['5101'].update(table='Products', child_key='Key', fields={'5102': 'Title'})
    result = convert_products_file('source.json', json.dumps({'Products': rows}).encode(), tmp_path, product_model_files=bundle(), product_mapping_profile=config)
    payload = json.loads((tmp_path / result['job_id'] / 'products.json').read_text(encoding='utf-8'))
    assert len(payload['products']) == 1
    assert result['report']['nested_relations'][0]['matched_records'] == 2


def test_separate_worksheet_source_survives_profile_save_and_imports_all_rows(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Documents'
    sheet.append(['Owner', 'Title', 'Lang'])
    for index in range(120):
        sheet.append(['P1', f'Document {index}', 'pl'])
    sheet.append(['P2', 'Certificate', 'en'])
    stream = BytesIO()
    workbook.save(stream)
    analysis = analyze_uploaded_file('documents.xlsx', stream.getvalue())
    config = profile()
    relation = config['_nested_relations']['5101']
    relation.update(source_mode='external', source={'filename': 'documents.xlsx', 'tables': analysis['tables']})
    relation['child_key'] = normalize_header('Owner')
    relation['fields'] = {'5102': normalize_header('Title'), '5103': normalize_header('Lang')}
    config = json.loads(json.dumps(config))
    result = convert_products_file('source.json', source(), tmp_path, product_model_files=bundle(), product_mapping_profile=config)
    assert result['report']['nested_relations'][0]['matched_records'] == 121
    payload = json.loads((tmp_path / result['job_id'] / 'products.json').read_text(encoding='utf-8'))
    attrs = [[a for a in p['dataVersions'][0]['productAttributes'] if a['ParentAttributeId'] == 5101] for p in payload['products']]
    assert [len(items) for items in attrs] == [240, 2, 0]
    assert len({a['RowI'] for a in attrs[0]}) == 120
